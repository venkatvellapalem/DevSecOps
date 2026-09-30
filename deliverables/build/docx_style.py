"""Document styling helpers shared by all four deliverables.

python-docx has no notion of paragraph shading or borders, so those need raw XML.
Keeping it here means the four documents cannot drift apart visually.
"""
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x11, 0x18, 0x27)
INK2 = RGBColor(0x37, 0x41, 0x51)
MUTED = RGBColor(0x6B, 0x72, 0x80)
DIM = RGBColor(0x9C, 0xA3, 0xAF)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GREEN = RGBColor(0x15, 0x80, 0x3D)
RED = RGBColor(0xC0, 0x23, 0x1C)
AMBER = RGBColor(0x9A, 0x46, 0x06)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FILL_SOFT = "F3F4F6"
FILL_BLUE = "EFF4FF"
FILL_GREEN = "ECFDF5"
FILL_RED = "FEF2F2"
FILL_AMBER = "FFFBEB"
FILL_CODE = "F5F6F8"

SANS = "Segoe UI"
MONO = "Consolas"


def shade(el, fill):
    """Apply a background fill to a paragraph or table cell."""
    pr = el._p.get_or_add_pPr() if hasattr(el, "_p") else el._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    pr.append(sh)


def left_bar(el, color="2563EB", size=18):
    """Thick coloured bar down the left edge of a paragraph — used for callouts."""
    pr = el._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    bd.append(left)
    pr.append(bd)


def bottom_rule(el, color="D1D5DB", size=6):
    pr = el._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), color)
    bd.append(b)
    pr.append(bd)


def keep_with_next(p):
    p.paragraph_format.keep_with_next = True


def para(doc, text="", size=10.5, color=INK2, bold=False, italic=False,
         before=0, after=6, align=None, indent=0, font=SANS, line=1.25):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent:
        pf.left_indent = Inches(indent)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = font
    return p


def rich(doc, parts, size=10.5, color=INK2, before=0, after=6, indent=0, align=None):
    """parts: list of (text, {bold/color/size/italic/font}) """
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.25
    if indent:
        pf.left_indent = Inches(indent)
    if align:
        p.alignment = align
    for t, o in parts:
        r = p.add_run(t)
        r.font.size = Pt(o.get("size", size))
        r.font.bold = o.get("bold", False)
        r.font.italic = o.get("italic", False)
        r.font.color.rgb = o.get("color", color)
        r.font.name = o.get("font", SANS)
    return p


def h1(doc, number, text, color=INK):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(20)
    pf.space_after = Pt(8)
    pf.keep_with_next = True
    if number:
        r = p.add_run(f"{number}  ")
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = BLUE
        r.font.name = SANS
    r = p.add_run(text)
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = SANS
    bottom_rule(p)
    return p


def h2(doc, text, color=INK, before=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = SANS
    return p


def h3(doc, text, color=INK2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = SANS
    return p


def bullet(doc, text, level=0, bold_lead=None, size=10.5, color=INK2):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.2
    pf.left_indent = Inches(0.28 + 0.24 * level)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = INK
        r.font.name = SANS
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = SANS
    return p


def numbered(doc, text, size=10.5, bold_lead=None, color=INK2):
    p = doc.add_paragraph(style="List Number")
    pf = p.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.2
    if bold_lead:
        r = p.add_run(bold_lead)
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = INK
        r.font.name = SANS
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = SANS
    return p


def code(doc, lines, fill=FILL_CODE):
    """A shaded monospace block. One paragraph per line keeps copy-paste usable."""
    if isinstance(lines, str):
        lines = lines.split("\n")
    for i, ln in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(6 if i == 0 else 0)
        pf.space_after = Pt(6 if i == len(lines) - 1 else 0)
        pf.line_spacing = 1.05
        pf.left_indent = Inches(0.12)
        shade(p, fill)
        r = p.add_run(ln if ln else " ")
        r.font.size = Pt(9)
        r.font.name = MONO
        r.font.color.rgb = INK2
    return None


def callout(doc, label, body, fill=FILL_BLUE, bar="2563EB", label_color=BLUE,
            body_color=INK2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.left_indent = Inches(0.1)
    shade(p, fill)
    left_bar(p, bar)
    r = p.add_run(label.upper())
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = label_color
    r.font.name = SANS

    p2 = doc.add_paragraph()
    pf2 = p2.paragraph_format
    pf2.space_before = Pt(0)
    pf2.space_after = Pt(9)
    pf2.line_spacing = 1.2
    pf2.left_indent = Inches(0.1)
    shade(p2, fill)
    left_bar(p2, bar)
    r = p2.add_run(body)
    r.font.size = Pt(10)
    r.font.color.rgb = body_color
    r.font.name = SANS


def table(doc, headers, rows, widths=None, size=9.5, header_fill="F3F4F6"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run(htxt)
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = INK
        r.font.name = SANS
        shade(hdr[i], header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if isinstance(val, tuple):
                txt, o = val
                r = p.add_run(txt)
                r.font.bold = o.get("bold", False)
                r.font.color.rgb = o.get("color", INK2)
            else:
                r = p.add_run(str(val))
                r.font.color.rgb = INK2
            r.font.size = Pt(size)
            r.font.name = SANS
    if widths:
        for i, wd in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(wd)
    para(doc, "", after=6)
    return t


def cover(doc, kind, title, subtitle, meta):
    para(doc, "", after=40)
    p = para(doc, kind.upper(), size=10, color=BLUE, bold=True, after=4)
    p.paragraph_format.space_before = Pt(10)
    para(doc, title, size=30, color=INK, bold=True, after=6, line=1.1)
    para(doc, subtitle, size=13, color=MUTED, after=22)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(16)
    bottom_rule(rule, "2563EB", 12)
    for k, v in meta:
        rich(doc, [(f"{k}   ", {"bold": True, "color": INK, "size": 10}),
                   (v, {"color": MUTED, "size": 10})], after=5)
    doc.add_page_break()


def toc(doc, items):
    h2(doc, "Contents", before=0)
    for num, title in items:
        rich(doc, [(f"{num}   ", {"bold": True, "color": BLUE, "size": 10.5}),
                   (title, {"color": INK2, "size": 10.5})], after=4, indent=0.1)
    doc.add_page_break()


def footer_pagenum(doc, label):
    """Page number + document label in the footer."""
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label + "   ·   page ")
        r.font.size = Pt(8)
        r.font.color.rgb = DIM
        r.font.name = SANS
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        p._p.append(fld)


def setup(doc, margins=0.85):
    for s in doc.sections:
        s.top_margin = Inches(margins)
        s.bottom_margin = Inches(margins)
        s.left_margin = Inches(margins)
        s.right_margin = Inches(margins)
    st = doc.styles["Normal"]
    st.font.name = SANS
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK2
