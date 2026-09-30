"""Build the student worksheet and its solutions document.

The solutions file mirrors the worksheet question-for-question so a marker can
work down both side by side.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Inches, Pt

from docx_style import (AMBER, BLUE, DIM, FILL_AMBER, FILL_BLUE, FILL_GREEN,
                        FILL_RED, GREEN, INK, INK2, MUTED, RED,
                        bullet, callout, code, cover, footer_pagenum, h1, h2, h3,
                        numbered, para, rich, setup, table, toc)

OUT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPO = "https://github.com/venkatvellapalem/DevSecOps"
DASH = "https://venkatvellapalem.github.io/DevSecOps/"


def new():
    d = Document()
    setup(d)
    return d


def answer_lines(d, n=3, label=None):
    """Ruled space for a handwritten answer."""
    for i in range(n):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run("_" * 104)
        r.font.size = Pt(9)
        r.font.color.rgb = DIM


def marks(d, text):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = BLUE


def q(d, number, text, m):
    rich(d, [(f"Q{number}  ", {"bold": True, "color": BLUE, "size": 10.5}),
             (text, {"color": INK, "size": 10.5}),
             (f"    [{m}]", {"color": MUTED, "size": 9})], after=4, before=8)


# ================================================================ WORKSHEET
def worksheet():
    d = new()
    cover(d, "Student worksheet", "Secure DevOps Pipeline",
          "Assess your understanding of the pipeline you built",
          [("Project", "BCSSL Cybersecurity Lab Series — Project 15"),
           ("Total marks", "100"),
           ("Time", "Approximately 90 minutes, plus the hands-on section"),
           ("Name", "_______________________________________"),
           ("Date", "_______________________________________")])

    h2(d, "How this is marked", before=0)
    table(d, ["Section", "Topic", "Marks"],
          [["A", "Concepts", "20"],
           ["B", "Reading the pipeline", "20"],
           ["C", "Hands-on: break a gate (practical)", "25"],
           ["D", "Analysis and reasoning", "20"],
           ["E", "Challenge: adapt the pipeline", "15"],
           ["", ("Total", {"bold": True}), ("100", {"bold": True})]],
          widths=[0.8, 4.6, 0.9])
    callout(d, "Before you start",
            "Sections A, B, D and E can be answered from the repository and the student guide. "
            "Section C requires you to actually run the pipeline and record what you observe. "
            "You will not be marked on whether the pipeline is green; you will be marked on whether "
            "you can explain what happened and why.", FILL_BLUE, "1D4ED8")

    # ---- A
    h1(d, "A", "Concepts")
    marks(d, "20 marks — 2 marks each")

    q(d, 1, "What do the letters S, A, S and T stand for, and what kind of thing does that gate "
            "inspect? What can it not see?", 2)
    answer_lines(d, 3)
    q(d, 2, "What does SCA inspect, and why can a project with flawless hand-written code still "
            "fail an SCA gate?", 2)
    answer_lines(d, 3)
    q(d, 3, "In run 2, the SCA gate passed while the image scan failed, on the same commit. "
            "Explain how that is possible and why it matters.", 2)
    answer_lines(d, 3)
    q(d, 4, "Why is gitleaks necessary when the pipeline already has a SAST gate and an SCA gate? "
            "Be specific about what the other two are blind to.", 2)
    answer_lines(d, 3)
    q(d, 5, "What is an image digest, and why should you deploy by digest rather than by the tag "
            "latest?", 2)
    answer_lines(d, 3)
    q(d, 6, "The pipeline sets ignore-unfixed: true on the image scan. Explain what this does and "
            "why omitting it makes the gate useless.", 2)
    answer_lines(d, 3)
    q(d, 7, "Explain the phrase a gate that is always red is worse than no gate at all.", 2)
    answer_lines(d, 3)
    q(d, 8, "What is the difference in purpose between the deploy job and the notify job? Why does "
            "notify use if: failure()?", 2)
    answer_lines(d, 3)
    q(d, 9, "The pipeline runs on GitHub's runners. Name two things the host machine does and two "
            "things it deliberately does not do.", 2)
    answer_lines(d, 3)
    q(d, 10, "Give one class of vulnerability that this pipeline cannot catch, and explain why the "
             "pipeline cannot catch it.", 2)
    answer_lines(d, 3)

    d.add_page_break()

    # ---- B
    h1(d, "B", "Reading the pipeline")
    marks(d, "20 marks")

    q(d, 11, "The deploy job begins with this line. In your own words, explain exactly what GitHub "
             "does with it, and what happens if one of the listed jobs fails.", 4)
    code(d, ["needs: [build, hadolint, sast, sca, gitleaks, image-scan]"])
    answer_lines(d, 4)

    q(d, 12, "Complete the table. For each gate, name the category of problem it is designed to "
             "catch, and give one concrete example of the kind of finding it produces.", 6)
    table(d, ["Gate", "Category of problem", "Example of a finding"],
          [["build", "", ""], ["hadolint", "", ""], ["sast", "", ""],
           ["sca", "", ""], ["gitleaks", "", ""], ["image-scan", "", ""]],
          widths=[1.1, 2.6, 2.6])

    q(d, 13, "Two gates in the original lab brief could never pass, no matter what the student did. "
             "Name the specific configuration setting responsible in each case, and explain the "
             "consequence.", 4)
    answer_lines(d, 4)

    q(d, 14, "The pipeline declares permissions: contents: read at the top of the workflow, and "
             "grants packages: write only on the deploy job. Explain why this is better than "
             "granting write permission to everything.", 3)
    answer_lines(d, 3)

    q(d, 15, "Explain the difference between the build job and the deploy job in this pipeline. Why "
             "is it a defect for the deploy job to rebuild the image from source?", 3)
    answer_lines(d, 4)

    d.add_page_break()

    # ---- C
    h1(d, "C", "Hands-on: break a gate")
    marks(d, "25 marks — this section requires you to run the pipeline")

    callout(d, "Safety rule",
            "Never push a real credential, key or password to test a secret scanner. Use a random "
            "high-entropy string that has no value. Work on a branch and delete it afterwards. "
            "Pushing a real secret to a public repository is a security incident, not an exercise.",
            FILL_RED, "B91C1C", RED)

    h3(d, "Task C1 — reproduce the failure (10 marks)")
    numbered(d, "Clone the repository and create a branch.")
    numbered(d, "Restore the vulnerable state (the student guide explains how).")
    numbered(d, "Push the branch and open a pull request.")
    numbered(d, "Watch the pipeline run and record the result for every gate below.")

    table(d, ["Gate", "Pass / Fail / Skipped", "What it reported (one line)"],
          [["build", "", ""], ["hadolint", "", ""], ["sast", "", ""],
           ["sca", "", ""], ["gitleaks", "", ""], ["image-scan", "", ""],
           ["deploy", "", ""], ["notify", "", ""]],
          widths=[1.1, 1.6, 3.6])

    h3(d, "Task C2 — plant a secret (8 marks)")
    numbered(d, "On the same branch, add a file containing a fabricated high-entropy string in a "
                "variable named something like API_SECRET.")
    numbered(d, "Push and record what the gitleaks gate reports.")
    numbered(d, "Explaining clearly why a real key would have been the wrong thing to use here.")

    answer_lines(d, 4)

    h3(d, "Task C3 — fix and verify (7 marks)")
    numbered(d, "Fix every failure you introduced.")
    numbered(d, "Push and confirm the pipeline goes green and deploy publishes.")
    numbered(d, "Record the image reference and digest from the run summary.")

    answer_lines(d, 4)

    d.add_page_break()

    # ---- D
    h1(d, "D", "Analysis and reasoning")
    marks(d, "20 marks")

    q(d, 16, "A colleague says: all three of these scanners overlap, so we should just run the "
             "cheapest one. Using evidence from the runs, construct the strongest possible argument "
             "against that position.", 5)
    answer_lines(d, 5)

    q(d, 17, "Your manager asks why the deploy job did not run on a pull request. Explain the "
             "design decision behind that, and whether you agree with it.", 4)
    answer_lines(d, 4)

    q(d, 18, "A dependency you use is flagged with a CVE that has no fix available from the "
             "maintainer. Your pipeline goes red. Describe at least two reasonable options and the "
             "trade-off of each.", 5)
    answer_lines(d, 5)

    q(d, 19, "Explain the provenance chain from a git commit to a running container. What does the "
             "digest prove, and what does it not prove?", 3)
    answer_lines(d, 4)

    q(d, 20, "The pipeline will eventually go red with no code change at all. Explain why, and "
             "describe how you would keep this from eroding trust in the gates.", 3)
    answer_lines(d, 3)

    d.add_page_break()

    # ---- E
    h1(d, "E", "Challenge")
    marks(d, "15 marks — open-ended")

    q(d, 21, "The pipeline's SAST and SCA gates are Python-specific. Describe exactly which parts of "
             "the workflow you would change to support a Node.js application, and which parts would "
             "need no change at all.", 5)
    answer_lines(d, 5)

    q(d, 22, "Propose a seventh gate for this pipeline. State what it would inspect, which tool you "
             "would use, what class of problem it would catch that the existing six do not, and how "
             "you would prove it works.", 6)
    answer_lines(d, 6)

    q(d, 23, "The gates currently block the deploy, but a reviewer can still merge code that fails "
             "them. Describe how you would close that gap, and discuss one downside of doing so.", 4)
    answer_lines(d, 5)

    para(d, "", after=10)
    p = d.add_paragraph()
    r = p.add_run("End of worksheet")
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = MUTED

    footer_pagenum(d, "Student worksheet")
    path = os.path.join(OUT, "project_student_worksheet.docx")
    d.save(path)
    print(f"  {os.path.basename(path)}  {os.path.getsize(path):,} bytes")


# =============================================================== SOLUTIONS
def solutions():
    d = new()
    cover(d, "Solutions", "Secure DevOps Pipeline",
          "Model answers for the student worksheet",
          [("Project", "BCSSL Cybersecurity Lab Series — Project 15"),
           ("Total marks", "100"),
           ("Note", "Answers are model answers, not the only correct answers."),
           ("Companion", "project_student_worksheet.docx")])

    callout(d, "How to use this document",
            "Section A has correct answers. Sections B to E are marked on reasoning rather than "
            "wording. Where a question is genuinely open, the model answer lists the points a good "
            "response would make rather than a single expected sentence.", FILL_BLUE, "1D4ED8")

    # ---- A
    h1(d, "A", "Concepts — model answers")
    marks(d, "20 marks")

    answers = [
        ("Static Application Security Testing. It inspects source code without executing it, "
         "looking for dangerous patterns such as shell=True, eval, or unsafe deserialisation. "
         "It cannot see anything in third-party libraries, because the flaw is not in code the "
         "team wrote. In this project it caught B602, a request parameter reaching a shell.", 2),
        ("Software Composition Analysis. It inspects the dependencies you consume, checking their "
         "versions against a database of known vulnerabilities. It can fail even when your own code "
         "is flawless, because you do not control the upstream project. Here it found 74 "
         "vulnerabilities across 8 packages.", 2),
        ("The two gates inspect different things. SCA only knows about packages listed in "
         "requirements.txt. Trivy inspects the finished image, including everything the base image "
         "brought with it. In run 2 the CVEs were in wheel and jaraco.context, which ship inside the "
         "Python base image and appear nowhere in the manifest. It matters because it proves the "
         "gates catch disjoint classes of problem, and that one scanner is not sufficient.", 2),
        ("SAST reads code patterns, SCA reads version numbers, and Trivy reads installed packages. "
         "None of them reads values. An AWS key sitting in a configuration file is perfectly valid "
         "code, depends on no vulnerable library, and installs nothing, so it passes all three "
         "silently. Gitleaks is the only gate that examines content for credentials.", 2),
        ("A digest is a cryptographic fingerprint of the exact image bytes. Two images sharing a "
         "digest are byte-identical. Tags such as latest are mutable labels that can be repointed, "
         "so pulling one later may give you different bytes than the ones that were scanned. "
         "Deploying by digest is what makes the claim the scanned artifact is the running artifact "
         "verifiable rather than hopeful.", 2),
        ("It restricts the gate to findings that have a fix available upstream. Without it, the "
         "44 unfixable operating-system CVEs in the base image make the gate red on every commit "
         "regardless of anything the team does. A gate that cannot be satisfied is bypassed or "
         "ignored, which is worse than having no gate.", 2),
        ("If a gate fails no matter what you do, people learn that it is noise and start routing "
         "around it. Once a team is in the habit of overriding gates, the gates that do matter lose "
         "their authority too. The value of a gate depends on it being both meaningful and "
         "satisfiable.", 2),
        ("deploy publishes the approved artifact; it is the thing being protected. notify exists to "
         "tell a human when something failed. It uses if: failure() so it is the exact inverse of "
         "the deploy job: silent while everything is healthy, and firing only when a release was "
         "blocked. Without it, a red run waits for someone to happen to notice.", 2),
        ("It does two things: docker pull to fetch the approved image, and docker run to start it. "
         "It deliberately never builds an image and never runs a scanner. All building and scanning "
         "happens on GitHub's ephemeral runners. This matters because it means the host cannot "
         "build an image that differs from the one that was scanned.", 2),
        ("Any of: runtime exploitation after a clean build; logic flaws such as broken access "
         "control that look fine to a pattern matcher; configuration drift between the repository "
         "and production; and problems introduced after the build, since the checks are a snapshot "
         "of a moment and the vulnerability database changes continuously.", 2),
    ]
    for i, (text, m) in enumerate(answers, 1):
        rich(d, [(f"Q{i}  ", {"bold": True, "color": BLUE, "size": 10.5}),
                 (text, {"size": 10.5})], after=8)

    # ---- B
    h1(d, "B", "Reading the pipeline — model answers")
    marks(d, "20 marks")

    rich(d, [("Q11  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("needs: declares dependencies between jobs. GitHub will not start the deploy job "
              "unless every job named in the list completed successfully. If any one of them "
              "failed, deploy is reported as skipped and never runs a single step. Crucially, "
              "nothing inside the deploy job re-checks anything -- the dependency graph itself is "
              "the security control.", {"size": 10.5})], after=8)

    rich(d, [("Q12  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("Model table (any reasonable example earns the mark):", {"size": 10.5})], after=4)
    table(d, ["Gate", "Category of problem", "Example finding"],
          [["build", "packaging", "image fails to build; dependency resolution error"],
           ["hadolint", "container definition", "DL3002 the last USER is root; DL3007 unpinned base image"],
           ["sast", "code you wrote", "B602 subprocess call with shell=True (command injection)"],
           ["sca", "dependencies you consume", "CVE in requests 2.19.1 leaking credentials on redirect"],
           ["gitleaks", "committed secrets", "a high-entropy API key committed to a source file"],
           ["image-scan", "the built artifact", "CVE in wheel or jaraco.context inside the base image"]],
          widths=[1.0, 2.1, 3.7])

    rich(d, [("Q13  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("First, the image scan used severity: 'CRITICAL,HIGH' with exit-code: '1' and no "
              "ignore-unfixed, so it failed on 44 unfixable operating-system CVEs forever. Second, "
              "the SAST gate used bandit -ll, which reports Medium severity and above, and the lab's "
              "own sample application always triggers B104 (binding to 0.0.0.0). The consequence in "
              "both cases is a permanently red gate, which trains people to bypass the pipeline.",
              {"size": 10.5})], after=8)

    rich(d, [("Q14  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("It is the principle of least privilege. The workflow as a whole needs only to read "
              "the repository, so it is granted only that. A single job -- deploy -- needs to write "
              "a package, so it opts in for itself alone. If any other job in the pipeline is "
              "compromised, for example through a malicious dependency, it holds a token that can "
              "only read. Granting write everywhere widens the blast radius of any single failure.",
              {"size": 10.5})], after=8)

    rich(d, [("Q15  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("build produces the image; deploy publishes it. They must operate on the same bytes. "
              "If the deploy job rebuilds from source, then the artifact scanned by the gates and "
              "the artifact pushed to the registry are two separate builds. Even from the same "
              "commit, a rebuild can pull a different base image or resolve a dependency "
              "differently. You would be scanning artifact A and shipping artifact B, which makes "
              "the entire scanning step decorative.", {"size": 10.5})], after=8)

    d.add_page_break()

    # ---- C
    h1(d, "C", "Hands-on — expected outcomes")
    marks(d, "25 marks — award marks for observation and explanation, not for the pipeline being green")

    h3(d, "Task C1 — expected result (10 marks)")
    table(d, ["Gate", "Expected", "Why"],
          [["build", "pass", "the image builds; the flaws are not build errors"],
           ["hadolint", "pass", "the sample Dockerfile is clean at warning level"],
           ["sast", "FAIL", "B602 -- subprocess call with shell=True in app.py"],
           ["sca", "FAIL", "dozens of CVEs from flask 2.0.1, requests 2.19.1, urllib3 1.23"],
           ["gitleaks", "pass", "no secrets in the restored state"],
           ["image-scan", "FAIL", "the same dependency CVEs visible inside the built image"],
           ["deploy", "skipped", "needs: held -- three gates failed, so the job never started"],
           ["notify", "runs", "if: failure() triggered by the failed gates"]],
          widths=[1.1, 1.0, 4.8])
    rich(d, [("Marking note.  ", {"bold": True, "color": INK}),
             ("The key observation is deploy = skipped, not deploy = failed. It did not run and "
              "fail; it never started. Full marks require the student to notice and state that "
              "distinction.", {"size": 10})], after=8)

    h3(d, "Task C2 — expected result (8 marks)")
    bullet(d, "gitleaks fails with a rule such as generic-api-key, and reports the file and line.")
    bullet(d, "Because it fails, deploy is skipped even if every other gate passed.")
    bullet(d, "The notify job fires, because if: failure() is now satisfied.")
    bullet(d, "The scanner redacts the value in its own output, so the CI log does not become a "
              "second copy of the secret. Students who mention redaction have understood a subtle "
              "point.")
    callout(d, "Why a fabricated string was required",
            "A real credential pushed to a public repository is compromised the moment it is "
            "pushed, and must be rotated immediately. The exercise needs a value with no real "
            "consequence. Note also that Gitleaks allowlists AWS's documentation example key, so "
            "testing with AKIAIOSFODNN7EXAMPLE produces a misleading pass.",
            FILL_AMBER, "D97706", AMBER)

    h3(d, "Task C3 — expected result (7 marks)")
    bullet(d, "sast fixed by removing the vulnerable endpoint or replacing shell=True with the "
              "argv form")
    bullet(d, "sca fixed by updating every pinned package in requirements.txt")
    bullet(d, "the image scan typically needs the base image build tooling removed or upgraded, "
              "not just the application dependencies")
    bullet(d, "all gates green, deploy succeeds, and the run summary shows a ghcr.io reference "
              "with a digest")
    rich(d, [("Marking note.  ", {"bold": True, "color": INK}),
             ("Students who expect the image scan to pass as soon as requirements.txt is updated "
              "have missed run 2. That is the single most common wrong answer and it should cost "
              "no more than one mark if the reasoning is otherwise sound.", {"size": 10})], after=8)

    d.add_page_break()

    # ---- D
    h1(d, "D", "Analysis — model answers")
    marks(d, "20 marks")

    rich(d, [("Q16  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("A strong answer uses evidence rather than assertion. Run 2 is the proof: the "
              "dependency gate passed while the image scan failed on the same commit, because "
              "Trivy detected CVEs in packages the dependency gate cannot see at all. After that, "
              "point out that the gates inspect four genuinely distinct inputs -- the container "
              "definition, the repository's own source, the dependency manifest, and the built "
              "artifact -- plus secrets, which are values rather than any of those. Choosing the "
              "cheapest scanner would mean choosing which class of vulnerability to ship. Full "
              "marks for naming run 2 explicitly.", {"size": 10.5})], after=8)

    rich(d, [("Q17  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("The deploy job is gated on github.ref == refs/heads/main and the event being a push, "
              "so a pull request runs the gates but publishes nothing. The reasoning is that a pull "
              "request is a proposal, and publishing an unreviewed image to the registry would "
              "defeat the point of review. A reasonable disagreement is that you might want to "
              "publish pull-request images under a branch tag for integration testing. Either "
              "position earns full marks if the reasoning is stated clearly.", {"size": 10.5})], after=8)

    rich(d, [("Q18  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("Expected options include: (a) add the finding to an ignore file with a recorded "
              "reason and expiry date, trading coverage for progress but requiring the decision to "
              "be visible and revisited; (b) accept the risk formally and document it, which is "
              "legitimate but must be a conscious decision rather than a silent override; "
              "(c) replace or vendor the dependency, the most correct option but sometimes not "
              "possible; (d) where the finding has no fix upstream, confirm that ignore-unfixed is "
              "already in place so only fixable findings gate the build. Full marks require two "
              "options and an honest trade-off for each.", {"size": 10.5})], after=8)

    rich(d, [("Q19  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("A git commit triggers the pipeline; the pipeline builds exactly one image; the gates "
              "scan that image; the deploy job publishes the same bytes and records a digest; the "
              "host pulls by that digest. The digest proves the artefact the host runs is "
              "byte-identical to the artefact that was scanned. It does not prove the image is "
              "secure, does not prove it was built from that specific commit unless the provenance "
              "chain is intact, and does not prove the scanner's database was complete.",
              {"size": 10.5})], after=8)

    rich(d, [("Q20  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("The vulnerability databases are updated continuously, so a CVE disclosed today "
              "against a package you already use will fail a commit that has not changed. Run 2 is "
              "the demonstration. To preserve trust: keep ignore-unfixed so the team only sees "
              "actionable findings, keep the failure notification so nobody is surprised, treat "
              "dependency updates as routine maintenance rather than incidents, and use Dependabot "
              "or similar so the fix arrives as a pull request. The important point is that a "
              "suddenly red pipeline with an unchanged commit is expected behaviour, not a broken "
              "pipeline.", {"size": 10.5})], after=8)

    d.add_page_break()

    # ---- E
    h1(d, "E", "Challenge — model answers")
    marks(d, "15 marks — open-ended; award for specificity and justification")

    rich(d, [("Q21  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("Change the sast job to use a multi-language scanner such as Semgrep, or swap Bandit "
              "for the ecosystem's equivalent. Change the sca job to run npm audit, or osv-scanner "
              "which covers many ecosystems, against package-lock.json instead of pip-audit against "
              "requirements.txt. Update the input defaults accordingly. Everything else needs no "
              "change at all: hadolint reads a Dockerfile regardless of language, Gitleaks scans "
              "text files and history regardless of language, Trivy scans any image, and the "
              "artifact-passing and needs: structure are entirely language-agnostic. Full marks for "
              "stating explicitly that four of six gates are unaffected.", {"size": 10.5})], after=8)

    rich(d, [("Q22  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("Any well-argued proposal. Strong candidates include: a licence-compliance scanner to "
              "catch a copyleft dependency being pulled into a proprietary product; a DAST scan "
              "against a running instance to catch runtime issues the static gates cannot; an IaC "
              "scanner such as Checkov or tfsec if the repository defines infrastructure; "
              "Docker build-provenance verification via cosign to prove the image was produced by "
              "this pipeline rather than merely that its digest matches; or dead-code and "
              "dependency-freshness checks. Marks are for naming what it inspects, which tool, "
              "which gap it closes, and how the student would demonstrate it failing. A proposal "
              "that duplicates an existing gate scores poorly however good the tool.",
              {"size": 10.5})], after=8)

    rich(d, [("Q23  ", {"bold": True, "color": BLUE, "size": 10.5}),
             ("Add branch protection on main, and mark the gates as required status checks so the "
              "pull request cannot be merged while they are red. The downside is that it removes "
              "the human override for genuine emergencies, and a flaky or newly red gate can then "
              "block unrelated work entirely. Mitigations worth mentioning: an explicit and "
              "audited bypass for administrators, keeping ignore-unfixed so there is little "
              "spurious redness, and making sure someone is notified the moment a gate goes red so "
              "it is fixed quickly rather than blocking everyone for a day.", {"size": 10.5})], after=8)

    para(d, "", after=12)
    p = d.add_paragraph()
    r = p.add_run("End of solutions")
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = MUTED

    footer_pagenum(d, "Worksheet solutions")
    path = os.path.join(OUT, "project_student_worksheet_solutions.docx")
    d.save(path)
    print(f"  {os.path.basename(path)}  {os.path.getsize(path):,} bytes")


if __name__ == "__main__":
    worksheet()
    solutions()
